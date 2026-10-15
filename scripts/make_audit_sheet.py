"""Build the independent fact-equivalence audit workbook (reviewer W2): a stratified sample of facts with both
renderings and the pinned-spec clause each is anchored to. Deterministic (seed 20261015)."""
import random, re, sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]; sys.path.insert(0, str(ROOT / "src"))
from agentisa.fact_inventory import FACTS
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

spec = (ROOT / "data/vendor/stablehlo_spec.md").read_text()


def spec_excerpt(anchor):
    a = anchor.split("#", 1)[1]
    m = re.match(r"([a-z_]+)(?: \(([C\d, -]+)\))?", a)
    op, labels = m.group(1), m.group(2)
    sec = spec.split(f"\n### {op}\n", 1)
    if len(sec) < 2:
        return f"[general rule from spec section '{a}', not a single numbered constraint]"
    body = sec[1].split("\n### ", 1)[0]
    if labels:
        nums = set()
        for part in labels.replace(" ", "").split(","):
            if "-" in part:
                lo, hi = part.replace("C", "").split("-"); nums |= set(range(int(lo), int(hi) + 1))
            elif part:
                nums.add(int(part.replace("C", "")))
        out, keep = [], False
        for l in (x.strip() for x in body.splitlines()):
            mm = re.match(r"\* \(C(\d+)\)", l)
            if mm: keep = int(mm.group(1)) in nums
            elif l.startswith("#"): keep = False
            if keep and l: out.append(l)
        return " ".join(out)[:900] or body[:600]
    sem = body.split("#### Semantics", 1)[-1].split("####", 1)[0]
    return " ".join(re.sub(r"!\[.*?\]\(.*?\)", "", sem).split())[:900]


rng = random.Random(20261015)
sample = []
for cat in ["syntax", "type", "shape", "semantics"]:
    by_op = {}
    for f in FACTS:
        if f["category"] == cat: by_op.setdefault(f["op"], []).append(f)
    ops = sorted(by_op); rng.shuffle(ops)
    sample += [rng.choice(by_op[o]) for o in ops[:6]]

wb = Workbook()
ins = wb.active; ins.title = "Instructions"
YEL = PatternFill("solid", fgColor="FFFF00"); HDR = PatternFill("solid", fgColor="D9D9D9")
F = lambda **k: Font(name="Arial", **k)
rows = [
    ("Independent fact-equivalence audit (AgentISA-Bench, Study 2)", True),
    ("Purpose: check that each fact's prose sentence (condition B) and formal predicate (condition C) say the same thing, and that both are faithful to the pinned StableHLO spec clause they cite.", False),
    (f"Sample: {len(sample)} of {len(FACTS)} facts, stratified: 6 per category (syntax, type, shape, semantics), one per randomly chosen operation group, seed 20261015.", False),
    ("Spec: openxla/stablehlo docs/spec.md at commit 23d2f926. The 'Spec clause' column is extracted from that file.", False),
    ("", False),
    ("What to fill in (yellow cells on the 'Audit' sheet only):", True),
    ("Q1  Prose faithful to spec?  Y = states nothing the spec contradicts and omits nothing the fact claims to state; N otherwise.", False),
    ("Q2  Formal faithful to spec?  Same criterion for the formal predicate.", False),
    ("Q3  Prose vs formal:  Equivalent | Prose says more | Formal says more | Conflict.", False),
    ("     'Says more' = one rendering states a condition, edge case or example that the other lacks.", False),
    ("Notes: describe any discrepancy precisely, quoting the words.", False),
    ("", False),
    ("Example of a completed row (illustrative only; not part of the sample and not counted):", True),
    ("Fact ELT.SHP.1 | Q1 = Y | Q2 = Y | Q3 = Prose says more | Notes: prose adds 'a vector or scalar is never combined with a matrix without an explicit broadcast_in_dim first'; the predicate shape(lhs)=shape(rhs)=shape(result) implies it but does not state it.", False),
    ("", False),
    ("Auditor: ____________________   Date: ____________   (The auditor should not have authored the facts.)", False),
]
for i, (t, b) in enumerate(rows, 1):
    c = ins.cell(row=i, column=1, value=t); c.font = F(bold=b, size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ins.column_dimensions["A"].width = 150

ws = wb.create_sheet("Audit")
hdr = ["#", "Fact ID", "Op group", "Category", "Spec anchor", "Spec clause (pinned spec)", "Prose rendering (B)",
       "Formal rendering (C)", "Q1 Prose faithful? (Y/N)", "Q2 Formal faithful? (Y/N)", "Q3 Prose vs formal", "Notes"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = F(bold=True); c.fill = HDR
    c.alignment = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF")
for i, f in enumerate(sample, 2):
    vals = [i - 1, f["id"], f["op"], f["category"], f["source_anchor"].split("#", 1)[1],
            spec_excerpt(f["source_anchor"]), f["nl"], f["formal"], None, None, None, None]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = F(size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if j >= 9: c.fill = YEL
n = len(sample) + 1
yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); ws.add_data_validation(yn); yn.add(f"I2:J{n}")
q3 = DataValidation(type="list", formula1='"Equivalent,Prose says more,Formal says more,Conflict"', allow_blank=True)
ws.add_data_validation(q3); q3.add(f"K2:K{n}")
for col, w in zip("ABCDEFGHIJKL", [4, 11, 14, 10, 20, 48, 48, 48, 11, 11, 16, 36]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C2"

sm = wb.create_sheet("Summary")
items = [("Facts in sample", f"=COUNTA(Audit!B2:B{n})"),
         ("Rows completed (Q3 answered)", f"=COUNTA(Audit!K2:K{n})"),
         ("Q1 prose faithful = Y", f'=COUNTIF(Audit!I2:I{n},"Y")'),
         ("Q2 formal faithful = Y", f'=COUNTIF(Audit!J2:J{n},"Y")'),
         ("Q3 Equivalent", f'=COUNTIF(Audit!K2:K{n},"Equivalent")'),
         ("Q3 Prose says more", f'=COUNTIF(Audit!K2:K{n},"Prose says more")'),
         ("Q3 Formal says more", f'=COUNTIF(Audit!K2:K{n},"Formal says more")'),
         ("Q3 Conflict", f'=COUNTIF(Audit!K2:K{n},"Conflict")'),
         ("Equivalent share of completed rows", "=IFERROR(B6/B3,0)")]
sm.cell(row=1, column=1, value="Summary (updates automatically from the Audit sheet)").font = F(bold=True, size=12)
for i, (k, v) in enumerate(items, 2):
    sm.cell(row=i, column=1, value=k).font = F(); sm.cell(row=i, column=2, value=v).font = F()
sm.cell(row=10, column=2).number_format = "0.0%"
sm.column_dimensions["A"].width = 40; sm.column_dimensions["B"].width = 14
out = ROOT / "audit/fact_equivalence_audit.xlsx"; out.parent.mkdir(exist_ok=True); wb.save(out)
print(out, len(sample), [f["id"] for f in sample])
